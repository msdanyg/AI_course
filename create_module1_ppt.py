#!/usr/bin/env python3
"""
Create Module 1 PowerPoint Presentation
The Cognitive Shift: Understanding the Reasoning Engine

Uses Official ActivTrak Brand Guidelines from slide_layouts.md
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ActivTrak Official Brand Colors (from slide_layouts.md)
DARK_NAVY = RGBColor(0x14, 0x20, 0x3F)      # #14203F - Body text, closing slide bg
MEDIUM_NAVY = RGBColor(0x19, 0x31, 0x6A)    # #19316A - Header bars, accents
BLUE = RGBColor(0x16, 0x57, 0xA0)           # #1657A0 - Section dividers
WHITE = RGBColor(0xFF, 0xFF, 0xFF)          # #FFFFFF - Background, text on dark
TEAL = RGBColor(0x2E, 0xD4, 0xB5)           # #2ED4B5 - CTAs, accent lines
YELLOW = RGBColor(0xFF, 0xC8, 0x3B)         # #FFC83B - Secondary accents (sparingly)

def set_slide_background(slide, color):
    """Set slide background color"""
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_title_slide(prs, title, subtitle):
    """Title Slide with white background and Dark Navy accent bar"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)

    # Dark Navy accent bar at top
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.5))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = DARK_NAVY
    accent_bar.line.fill.background()

    # Title - centered
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(9), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = DARK_NAVY
    p.font.name = "Century Gothic"
    p.alignment = PP_ALIGN.CENTER

    # Subtitle - centered
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.8))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(24)
        p.font.color.rgb = MEDIUM_NAVY
        p.font.name = "Century Gothic"
        p.alignment = PP_ALIGN.CENTER

    return slide

def add_section_slide(prs, title):
    """Section Divider with Blue background and Teal accent"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, BLUE)

    # Teal accent line
    accent_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3), Inches(3), Inches(4), Inches(0.05))
    accent_line.fill.solid()
    accent_line.fill.fore_color.rgb = TEAL
    accent_line.line.fill.background()

    # Centered title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.3), Inches(9), Inches(1.2))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"
    p.alignment = PP_ALIGN.CENTER

    return slide

def add_content_slide(prs, title, bullets):
    """Content Slide with Medium Navy header bar"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)

    # Medium Navy header bar
    header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = MEDIUM_NAVY
    header_bar.line.fill.background()

    # Left-aligned title in header
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"

    # Bullets (20pt, Dark Navy)
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(9), Inches(5.3))
    tf = content_box.text_frame
    tf.word_wrap = True

    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {bullet}" if bullet else ""
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_NAVY
        p.font.name = "Century Gothic"
        p.space_after = Pt(6)

    return slide

def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    """Two-Column Slide for comparisons"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)

    # Medium Navy header bar
    header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = MEDIUM_NAVY
    header_bar.line.fill.background()

    # Left-aligned title in header
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"

    # Left column header
    left_header = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.2), Inches(0.5))
    tf = left_header.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = MEDIUM_NAVY
    p.font.name = "Century Gothic"

    # Left column content
    left_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.1), Inches(4.2), Inches(5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(left_items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_NAVY
        p.font.name = "Century Gothic"
        p.space_after = Pt(6)

    # Right column header
    right_header = slide.shapes.add_textbox(Inches(5.3), Inches(1.5), Inches(4.2), Inches(0.5))
    tf = right_header.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = MEDIUM_NAVY
    p.font.name = "Century Gothic"

    # Right column content
    right_box = slide.shapes.add_textbox(Inches(5.3), Inches(2.1), Inches(4.2), Inches(5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(right_items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_NAVY
        p.font.name = "Century Gothic"
        p.space_after = Pt(6)

    return slide

def add_closing_slide(prs, message, cta=""):
    """Closing Slide with Dark Navy background and Teal accent"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, DARK_NAVY)

    # Teal accent bar at bottom
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7), Inches(10), Inches(0.5))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = TEAL
    accent_bar.line.fill.background()

    # Centered closing message
    msg_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(9), Inches(1.2))
    tf = msg_box.text_frame
    p = tf.paragraphs[0]
    p.text = message
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"
    p.alignment = PP_ALIGN.CENTER

    # Optional CTA
    if cta:
        cta_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(0.8))
        tf = cta_box.text_frame
        p = tf.paragraphs[0]
        p.text = cta
        p.font.size = Pt(24)
        p.font.color.rgb = TEAL
        p.font.name = "Century Gothic"
        p.alignment = PP_ALIGN.CENTER

    return slide

def add_quote_slide(prs, quote, attribution=""):
    """Quote slide with Blue background"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, BLUE)

    # Quote
    quote_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2.5))
    tf = quote_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f'"{quote}"'
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"
    p.alignment = PP_ALIGN.CENTER

    if attribution:
        attr_box = slide.shapes.add_textbox(Inches(1), Inches(5.2), Inches(8), Inches(0.5))
        tf = attr_box.text_frame
        p = tf.paragraphs[0]
        p.text = attribution
        p.font.size = Pt(18)
        p.font.color.rgb = TEAL
        p.font.name = "Century Gothic"
        p.alignment = PP_ALIGN.CENTER

    return slide

def add_key_takeaways_slide(prs, takeaways):
    """Key takeaways slide with Blue header"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)

    # Blue header bar for emphasis
    header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = BLUE
    header_bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Key Takeaways"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"

    # Takeaways
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(9), Inches(5.3))
    tf = content_box.text_frame
    tf.word_wrap = True

    for i, takeaway in enumerate(takeaways):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = takeaway
        p.font.size = Pt(20)
        p.font.color.rgb = DARK_NAVY
        p.font.name = "Century Gothic"
        p.space_after = Pt(12)

    return slide

def add_water_glass_slide(prs):
    """Special slide for the Water Glass Effect illustration"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    set_slide_background(slide, WHITE)

    # Medium Navy header bar
    header_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = MEDIUM_NAVY
    header_bar.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "The Water Glass Effect"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Century Gothic"

    # Three glass representations
    glasses = [
        ("Paperclip", "Small displacement", TEAL, "MOST WATER"),
        ("Scissors", "Medium displacement", MEDIUM_NAVY, ""),
        ("Pocket Watch", "Large displacement", DARK_NAVY, "LEAST WATER")
    ]

    for i, (obj, desc, color, label) in enumerate(glasses):
        x = Inches(1.0 + (i * 2.8))

        # Glass shape (rectangle as placeholder)
        glass = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2.0), Inches(2.0), Inches(2.8))
        glass.fill.solid()
        glass.fill.fore_color.rgb = RGBColor(0xE0, 0xF2, 0xFE)  # Light blue for water
        glass.line.color.rgb = MEDIUM_NAVY
        glass.line.width = Pt(2)

        # Object label
        obj_box = slide.shapes.add_textbox(x, Inches(3.2), Inches(2.0), Inches(0.5))
        tf = obj_box.text_frame
        p = tf.paragraphs[0]
        p.text = obj
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color
        p.font.name = "Century Gothic"
        p.alignment = PP_ALIGN.CENTER

        # Description
        desc_box = slide.shapes.add_textbox(x, Inches(5.0), Inches(2.0), Inches(0.5))
        tf = desc_box.text_frame
        p = tf.paragraphs[0]
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = DARK_NAVY
        p.font.name = "Century Gothic"
        p.alignment = PP_ALIGN.CENTER

        # Label if winner
        if label:
            label_box = slide.shapes.add_textbox(x, Inches(5.5), Inches(2.0), Inches(0.5))
            tf = label_box.text_frame
            p = tf.paragraphs[0]
            p.text = label
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = TEAL
            p.font.name = "Century Gothic"
            p.alignment = PP_ALIGN.CENTER

    # Lesson callout
    lesson_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(6.2), Inches(9), Inches(0.8))
    lesson_box.fill.solid()
    lesson_box.fill.fore_color.rgb = RGBColor(0xF0, 0xF8, 0xF5)
    lesson_box.line.color.rgb = TEAL
    lesson_box.line.width = Pt(2)

    lesson_text = slide.shapes.add_textbox(Inches(0.7), Inches(6.35), Inches(8.6), Inches(0.5))
    tf = lesson_text.text_frame
    p = tf.paragraphs[0]
    p.text = "Irrelevant details scatter AI attention away from the core problem"
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_NAVY
    p.font.name = "Century Gothic"
    p.alignment = PP_ALIGN.CENTER

    return slide

def create_presentation():
    """Create the full Module 1 presentation with ActivTrak branding"""
    prs = Presentation()
    # Official dimensions: 10" x 7.5" (4:3 aspect ratio)
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title
    add_title_slide(prs,
        "Module 1: The Cognitive Shift",
        "Understanding the Reasoning Engine")

    # Slide 2: Opening Hook
    add_quote_slide(prs,
        "Someone pastes a giant email thread into Claude. Types 'what should I do?' The response is completely useless.",
        "This isn't the AI being dumb. It's a mental model mismatch.")

    # Slide 3: Section - Not a Search Engine
    add_section_slide(prs, "Not a Search Engine")

    # Slide 4: Search vs Reasoning
    add_two_column_slide(prs, "Search Engine vs Reasoning Engine",
        "Search Engine (Google)",
        ["Retrieves existing content", "More keywords = better results", "Finds pages containing words", "Returns links to sources"],
        "Reasoning Engine (Claude)",
        ["Generates new content", "More words = potential confusion", "Predicts probable next words", "Creates original responses"])

    # Slide 5: The Problem
    add_content_slide(prs, "The Problem with Dumping Everything", [
        "When you dump everything into a prompt, you're not giving AI 'more to work with'",
        "You're giving it more possible paths to wander down - most of them wrong",
        "Technical term: Probabilistic reasoning engine",
        "Your job: Constrain probabilities so high-probability paths lead where you want"
    ])

    # Slide 6: Section - Water Glass Effect
    add_section_slide(prs, "The Water Glass Effect")

    # Slide 7: Water Glass Visual
    add_water_glass_slide(prs)

    # Slide 8: Water Glass in Practice
    add_content_slide(prs, "Water Glass Effect in Your Work", [
        "Pasting a 20-message email thread and asking for a 'strategic recommendation'",
        "Claude's attention scatters across: holiday party logistics, personnel gossip, emotional subtext AND the business question",
        "The response addresses all of these poorly instead of any of them well",
        "The fix: Explicit focus - tell Claude what to focus on AND what to ignore"
    ])

    # Slide 9: The Fix
    add_content_slide(prs, "Concentrate AI Attention", [
        "Bad: 'Give me a strategic recommendation based on this email thread'",
        "Good: 'Identify the three highest-priority roadmap items. Ignore discussion of social events or personnel changes.'",
        "Tell the AI what to focus on",
        "Tell the AI what to ignore",
        "Curate ruthlessly"
    ])

    # Slide 10: Section - Solo Pilot vs Squadron Leader
    add_section_slide(prs, "Solo Pilot vs Squadron Leader")

    # Slide 11: Solo Pilot Problems
    add_content_slide(prs, "Solo Pilot Problems", [
        "Context bleeds everywhere - instructions mix with data, background contaminates the task",
        "Errors cascade unchecked - misunderstand step one, everything after builds on that mistake",
        "No specialization - one 'generalist' AI tries to research, strategize, write AND critique simultaneously",
        "It does all of them poorly"
    ])

    # Slide 12: Squadron Leader Approach
    add_content_slide(prs, "Squadron Leader Approach", [
        "Break complex tasks into specialized stages",
        "Feed structured context appropriate to each stage",
        "Verify before proceeding to next stage",
        "Each agent excels at its function",
        "Coordinate outputs into something none could produce alone"
    ])

    # Slide 13: Comparison
    add_two_column_slide(prs, "The Difference",
        "Solo Pilot",
        ["One prompt, hope for the best", "Results vary wildly", "Individual skill dependent", "Paste everything, start over"],
        "Squadron Leader",
        ["Coordinated workflow", "Consistent quality", "Process-encoded expertise", "Structured handoffs"])

    # Slide 14: Section - Context as Mission Briefing
    add_section_slide(prs, "Context as Mission Briefing")

    # Slide 15: Context Window
    add_content_slide(prs, "The Context Window", [
        "Claude's context window: About 500 pages worth of information",
        "The trap: Thinking more context equals better results",
        "Research shows: Information placement matters enormously",
        "Beginning and end get high attention - the middle gets lost"
    ])

    # Slide 16: Good Prompt Structure
    add_content_slide(prs, "Structure Like a Mission Briefing", [
        "Opening orientation: Who you are, what we're doing",
        "Core data: The specific information needed",
        "Clear instructions: What to do with that data",
        "Output specs: Format and constraints",
        "Same information, radically different results"
    ])

    # Slide 17: Section - ActivTrak Connection
    add_section_slide(prs, "Insights, Not Oversight")

    # Slide 18: Framing Matters
    add_two_column_slide(prs, "How You Frame the Task Matters",
        "Oversight Framing",
        ["'Analyze this data to find who's underperforming'", "Surveillance-style outputs", "Punitive focus", "Defensive reactions"],
        "Insight Framing",
        ["'Identify teams that might need additional support'", "Consultative insights", "Supportive focus", "Constructive outcomes"])

    # Slide 19: Key Takeaways
    add_key_takeaways_slide(prs, [
        "1. AI is a reasoning engine, not a search engine - structure your context, don't dump keywords",
        "2. The Water Glass Effect is real - irrelevant context degrades AI reasoning. Curate ruthlessly.",
        "3. Solo Pilots improvise. Squadron Leaders orchestrate. Complex work requires specialized agents and structured handoffs."
    ])

    # Slide 20: What's Next
    add_content_slide(prs, "What's Next?", [
        "Lab Exercise: Experience the Water Glass Effect firsthand",
        "Practice restructuring prompts that actually work",
        "Module 2: Model Selection",
        "When to use Claude's different thinking modes",
        "When Gemini is the right tool for the job"
    ])

    # Slide 21: Closing
    add_closing_slide(prs,
        "The Cognitive Shift Starts Now",
        "See you in the lab.")

    # Save
    output_path = "/Users/dglickman@bgrove.com/AI course/Module 1 - The Cognitive Shift/Module_1_Slides.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    create_presentation()
